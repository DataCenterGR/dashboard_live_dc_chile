import openpyxl

path = r'c:\Users\CristianJoseRamosCon\OneDrive - Grenergy Renovables\Escritorio\DC Code\20260901 Analisis estados de proyecto GR (1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print('SHEETS', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\n### {ws.title}')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
        print(row)
